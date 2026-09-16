"""VLM OCR 回归（需求 F1.9 / M2）：PDF 红页整页 OCR + PDF/DOCX 内嵌图 OCR + 版面对齐。

覆盖三条链路：

- 红页（PDF）：图片型扫描页（无文本层 → 判 red）→ 整页 OCR → source=vlm_ocr 块；
- PDF 内嵌图：混排页（文本层 + 嵌入带字图）→ 抽图 OCR → source=vlm_ocr_image 块，
  并**按图的纵坐标锚定插入位**（插到图上方最近文本块之后，实现图文版面对齐）；
- DOCX 内嵌图：段落内联图 → 取关系部件字节 → OCR → **原位替换** [图片] 占位块（段落顺序对齐）。

离线（无需 Key）用 FakeTranscriber 验编排；真实（VLM_ENABLED=1 + Key）用 Qwen-VL 验还原文字。

关键注意：pipeline 用的是「本模块导入的 get_transcriber」，故离线分支必须 patch
`app.ingestion.pipeline.get_transcriber`（patch vlm 模块同名属性不生效）。

退出码 0=全过 / 1=有失败。
"""
from __future__ import annotations

import io
import sys
import tempfile
from contextlib import contextmanager
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import pymupdf
from PIL import Image, ImageDraw, ImageFont

from app.core.config import settings
from app.ingestion.parsers import parse_file
from app.ingestion.pipeline import IngestPipeline
from app.ingestion.vlm import (
    DashScopeTranscriber,
    Transcriber,
    get_transcriber,
)
from app.models import Block

PASS: list[str] = []
FAIL: list[tuple[str, str]] = []


def check(name: str, cond: bool, detail: str = "") -> None:
    if cond:
        PASS.append(name)
        print(f"✓ {name}")
    else:
        FAIL.append((name, detail))
        print(f"✗ {name}  ({detail})")


def _draw_text_png(lines: list[str], w: int, h: int, font_px: int) -> bytes:
    """把文字烤进 PNG（模拟扫描件 / 带字示意图，无文本层）。"""
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("C:/Windows/Fonts/msyh.ttc", font_px)
    except Exception:
        font = ImageFont.load_default()
    y = max(20, font_px)
    for ln in lines:
        d.text((60, y), ln, fill="black", font=font)
        y += int(font_px * 1.6)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _make_scan_pdf(png_bytes: bytes, tmp_dir: str) -> Path:
    """整页只放一张图（无文本层 → PyMuPDF 抽不到字 → 判 red）。"""
    tmp = Path(tmp_dir) / "vlm_scan_fixture.pdf"
    doc = pymupdf.open()
    page = doc.new_page(width=1240, height=1754)
    page.insert_image(page.rect, stream=png_bytes)
    doc.save(str(tmp))
    doc.close()
    return tmp


def _make_mixed_pdf(tmp_dir: str, *, with_below: bool = True) -> Path:
    """图文混排页：上/下文本层（>100 字符 → 非 red）+ 中间一张带字内嵌图。

    上下文本各带 ASCII 标记（ABOVEMARKER/BELOWMARKER），便于校验 OCR 块被插到图位置。
    """
    tmp = Path(tmp_dir) / "vlm_mixed_fixture.pdf"
    figure = _draw_text_png(["设备保养记录", "办公设备每季度保养一次。"], 800, 320, 40)

    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)  # A4
    above = (
        "ABOVEMARKER This paragraph is above the embedded figure and keeps the "
        "page grade non-red because it carries a normal text layer of its own."
    )
    page.insert_textbox(pymupdf.Rect(50, 50, 545, 150), above, fontsize=11)
    page.insert_image(pymupdf.Rect(50, 300, 545, 460), stream=figure)
    if with_below:
        below = "BELOWMARKER This paragraph is below the embedded figure."
        page.insert_textbox(pymupdf.Rect(50, 520, 545, 600), below, fontsize=11)
    doc.save(str(tmp))
    doc.close()
    return tmp


def _make_docx_with_image(tmp_dir: str) -> Path:
    """DOCX：标题 + 说明段 + 纯内联图段（带字）+ 结尾段，用于验证占位块原位替换。"""
    from docx import Document
    from docx.shared import Inches

    tmp = Path(tmp_dir) / "vlm_docx_fixture.docx"
    figure = _draw_text_png(["设备保养记录", "办公设备每季度保养一次。"], 800, 320, 40)
    doc = Document()
    doc.add_heading("设备保养说明", level=1)
    doc.add_paragraph("下方图片为保养周期示意。")
    doc.add_picture(io.BytesIO(figure), width=Inches(4))
    doc.add_paragraph("以上为全部内容。")
    doc.save(str(tmp))
    return tmp


class _FakeTranscriber(Transcriber):
    def __init__(self, text: str):
        self._text = text
        self.calls = 0
        self.model = "fake"
        self.max_pages = 20  # _transcribe_red_pages 会读该属性

    def transcribe(self, image_bytes: bytes, page_no: int, *,
                   mime: str = "image/png", prompt: str | None = None) -> str | None:
        self.calls += 1
        return self._text


@contextmanager
def _patched_transcriber(fake: Transcriber):
    """patch pipeline 实际使用的 get_transcriber（而非 vlm 模块同名属性）。"""
    import app.ingestion.pipeline as pl_mod

    orig = pl_mod.get_transcriber
    pl_mod.get_transcriber = lambda: fake
    try:
        yield
    finally:
        pl_mod.get_transcriber = orig


def _offline_checks() -> None:
    print("===== 离线：编排逻辑（FakeTranscriber）=====")
    tmp_dir = tempfile.mkdtemp(prefix="vlm_ocr_")
    pl = IngestPipeline()

    # ── 红页路径（PDF）──
    png = _draw_text_png([
        "员工请假管理制度（扫描件）",
        "第一条 员工每年享有带薪年假 5 天。",
        "第二条 病假须提供正规医院出具的诊断证明。",
    ], 1240, 1754, 40)
    red_fixture = _make_scan_pdf(png, tmp_dir)
    parsed = parse_file(red_fixture)
    red = [q.page_no for q in parsed.page_qualities if q.level == "red"]
    check("扫描 fixture 被判 red（0 文本字符）", bool(red),
          f"quals={[(q.level, q.text_chars) for q in parsed.page_qualities]}")

    with _patched_transcriber(_FakeTranscriber("【OCR】员工每年享有带薪年假 5 天。")):
        blocks = pl._transcribe_red_pages(red_fixture, parsed, red)
    check("_transcribe_red_pages 产出 1 个 OCR 块", len(blocks) == 1, f"n={len(blocks)}")
    if blocks:
        b = blocks[0]
        check("OCR 块 page 溯源 = 红页页码", b.page == red[0], f"page={b.page} expect={red[0]}")
        check("OCR 块 metadata.source=vlm_ocr", b.metadata.get("source") == "vlm_ocr")
        check("OCR 块为 paragraph", b.block_type == "paragraph")

    # ── PDF 内嵌图路径 + 版面对齐 ──
    mixed = _make_mixed_pdf(tmp_dir)
    mp = parse_file(mixed)
    mred = [q.page_no for q in mp.page_qualities if q.level == "red"]
    check("混排 fixture 非 red（有文本层）", not mred,
          f"quals={[(q.level, q.text_chars) for q in mp.page_qualities]}")

    fig_fake = _FakeTranscriber("【图内文字】办公设备每季度保养一次。")
    with _patched_transcriber(fig_fake):
        inline = pl._transcribe_inline_images(mixed, mp, mred)
    n_inline = sum(len(v) for v in inline.values())
    check("_transcribe_inline_images 抽到 1 个内嵌图块", n_inline == 1, f"n={n_inline}")
    check("内嵌图走了 VLM（fake 被调用）", fig_fake.calls == 1, f"calls={fig_fake.calls}")
    anns = [a for v in inline.values() for a in v]  # [(pos, Block)]
    if anns:
        pos, b = anns[0]
        check("内嵌图块 page 溯源正确", b.page == 1, f"page={b.page}")
        check("内嵌图块 source=vlm_ocr_image", b.metadata.get("source") == "vlm_ocr_image")
        check("内嵌图块为 paragraph", b.block_type == "paragraph")
        # 版面对齐：插入位应落在「上方段」之后（页内块序 0-based，位序=1 → 插在第 2 块前）
        check("版面对齐：插入位在「上方段」之后（pos=1）", pos == 1, f"pos={pos}")

    # 合并后顺序：上方段 → 图 OCR → 下方段
    sim = parse_file(mixed)
    sim.blocks.append(Block(block_type="paragraph", text="鈭涗鈥稨亂碼",
                            page=1, metadata={"quality_level": "red"}))
    fake_annots = {1: [(1, Block(block_type="paragraph", text="图内文字", page=1,
                                 metadata={"source": "vlm_ocr_image"}))]}
    pl._merge_vlm_text_blocks(sim, {}, fake_annots)
    texts = [b.text for b in sim.blocks]
    i_above = next((i for i, t in enumerate(texts) if "ABOVEMARKER" in t), -1)
    i_ocr = next((i for i, t in enumerate(texts) if t == "图内文字"), -1)
    i_below = next((i for i, t in enumerate(texts) if "BELOWMARKER" in t), -1)
    check("合并后红页原文块已移除",
          not any(b.metadata.get("quality_level") == "red" for b in sim.blocks))
    check("版面对齐：合并后顺序 上方段 < 图OCR < 下方段",
          -1 < i_above < i_ocr < i_below, f"above={i_above} ocr={i_ocr} below={i_below}")

    # ── DOCX 内嵌图路径 ──
    dx = _make_docx_with_image(tmp_dir)
    dp = parse_file(dx)
    img_blocks = [b for b in dp.blocks if b.metadata.get("has_inline_image")]
    check("DOCX 解析出内嵌图占位块（带 rId）",
          len(img_blocks) == 1 and bool(img_blocks[0].metadata.get("img_rid")),
          f"n={len(img_blocks)}")
    with _patched_transcriber(_FakeTranscriber("【图内文字】办公设备每季度保养一次。")):
        n_docx = pl._transcribe_docx_images(dx, dp)
    check("DOCX 内嵌图 OCR 计数=1", n_docx == 1, f"n={n_docx}")
    dtexts = [b.text for b in dp.blocks]
    check("DOCX 占位块 [图片] 已被替换", "[图片]" not in dtexts, f"texts={dtexts}")
    check("DOCX OCR 块 source=vlm_ocr_docx",
          any(b.metadata.get("source") == "vlm_ocr_docx" for b in dp.blocks))
    i_ocr = next((i for i, b in enumerate(dp.blocks)
                  if b.metadata.get("source") == "vlm_ocr_docx"), -1)
    i_before = next((i for i, t in enumerate(dtexts) if "保养周期示意" in t), -1)
    i_after = next((i for i, t in enumerate(dtexts) if "以上为全部内容" in t), -1)
    check("DOCX OCR 块位置对齐（原图占位处）",
          -1 < i_before < i_ocr < i_after, f"before={i_before} ocr={i_ocr} after={i_after}")

    for f in (red_fixture, mixed, dx):
        Path(f).unlink(missing_ok=True)


def _real_vlm_checks() -> None:
    print("===== 真实：Qwen-VL（gate: VLM_ENABLED + Key）=====")
    if not (settings.vlm_enabled and settings.has_api_key):
        print("(skipped: VLM_ENABLED=0 或无 Key → 仅离线逻辑已验证)")
        return
    tr = get_transcriber()
    if not isinstance(tr, DashScopeTranscriber):
        print("(skipped: transcriber 非 DashScopeTranscriber)")
        return

    tmp_dir = tempfile.mkdtemp(prefix="vlm_ocr_real_")
    pl = IngestPipeline()

    # 红页整页 OCR
    png = _draw_text_png([
        "设备保养记录（扫描件）",
        "第一条 办公设备每季度保养一次。",
        "第二条 保养记录须归档保存两年。",
    ], 1240, 1754, 40)
    red_fixture = _make_scan_pdf(png, tmp_dir)
    parsed = parse_file(red_fixture)
    red = [q.page_no for q in parsed.page_qualities if q.level == "red"]
    blocks = pl._transcribe_red_pages(red_fixture, parsed, red)
    check("真实 VLM 红页转录产出 OCR 块", len(blocks) == 1, f"n={len(blocks)}")
    if blocks:
        txt = blocks[0].text or ""
        check("真实 VLM 红页恢复关键文字（保养/每季度）",
              "保养" in txt and "每季度" in txt, f"len={len(txt)}")

    # PDF 内嵌图 OCR + 版面对齐
    mixed = _make_mixed_pdf(tmp_dir)
    mp = parse_file(mixed)
    mred = [q.page_no for q in mp.page_qualities if q.level == "red"]
    inline = pl._transcribe_inline_images(mixed, mp, mred)
    anns = [a for v in inline.values() for a in v]
    inline_txt = "\n".join(b.text for _, b in anns)
    check("真实 VLM PDF 内嵌图 OCR 产出块", bool(inline_txt.strip()), f"len={len(inline_txt)}")
    check("真实 VLM PDF 内嵌图恢复关键文字（保养/每季度）",
          "保养" in inline_txt and "每季度" in inline_txt, f"txt={inline_txt[:80]!r}")
    if anns:
        check("真实 VLM PDF 内嵌图插入位对齐（pos=1）", anns[0][0] == 1, f"pos={anns[0][0]}")

    # DOCX 内嵌图 OCR
    dx = _make_docx_with_image(tmp_dir)
    dp = parse_file(dx)
    n_docx = pl._transcribe_docx_images(dx, dp)
    joined = "\n".join(b.text for b in dp.blocks)
    check("真实 VLM DOCX 内嵌图 OCR 产出块", n_docx >= 1, f"n={n_docx}")
    check("真实 VLM DOCX 恢复关键文字（保养/每季度）",
          "保养" in joined and "每季度" in joined, f"texts={joined[:80]!r}")

    for f in (red_fixture, mixed, dx):
        Path(f).unlink(missing_ok=True)


def main() -> int:
    _offline_checks()
    _real_vlm_checks()
    print(f"\n结果：PASS {len(PASS)} / FAIL {len(FAIL)}")
    if FAIL:
        for n, d in FAIL:
            print(f"  FAIL {n}: {d}")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
